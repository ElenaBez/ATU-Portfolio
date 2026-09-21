"""
Reads one ATU project summary workbook into structured pandas data.

Expected tabs (names matched case-insensitively, partial match allowed):
  - Budget and forecast
  - Reporting periods
  - Staff salaries and allocation
  - Actual travel and overhead cost
"""
from __future__ import annotations

import calendar
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

# --------------------------------------------------------------------------
# Category mapping: sheet label -> dashboard label (None = skip)
# --------------------------------------------------------------------------
CATEGORY_ORDER = ["Staff", "Overheads", "Travel & Subsistence", "Equipment", "EE & Service Costs"]


def map_category(label: str) -> str | None:
    s = str(label).strip().lower()
    if not s:
        return None
    if "total" in s or "infrastructure" in s:
        return None
    if "salar" in s or "staff" in s:
        return "Staff"
    if "overhead" in s:
        return "Overheads"
    if "travel" in s:
        return "Travel & Subsistence"
    if "equipment" in s:
        return "Equipment"
    if "service" in s or s.startswith("ee"):
        return "EE & Service Costs"
    return None


ACTUAL_WORDS = ("actual", "claim")          # "Actual", "Actual Spend", "Claim Approved"
COMPLETED_STATUSES = ("verified", "draft")  # statuses that mark a reported period
PERIOD_RE = re.compile(r"^\s*P(\d{1,3})\s*$", re.I)
CODE_RE = re.compile(r"^\s*P\d{4,}\s*$", re.I)


# --------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------
def to_num(v) -> float:
    """Numbers stay numbers; dashes, blanks and text become 0."""
    if v is None:
        return 0.0
    if isinstance(v, bool):
        return float(v)
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "").replace("€", "")
    if s in ("", "-", "–", "—"):
        return 0.0
    try:
        return float(s)
    except ValueError:
        return 0.0


def is_blank(v) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "")


def to_date(v, warnings: list[str] | None = None, where: str = "") -> date | None:
    """Accepts datetime, Excel serial or dd/mm/yyyy text. Invalid days
    (e.g. 31/11/2025) are clamped to the last day of that month."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)):
        return (datetime(1899, 12, 30) + timedelta(days=float(v))).date()
    m = re.match(r"^\s*(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{2,4})\s*$", str(v))
    if m:
        d, mo, y = (int(x) for x in m.groups())
        if y < 100:
            y += 2000
        last = calendar.monthrange(y, mo)[1]
        if d > last:
            if warnings is not None:
                warnings.append(f"{where}: '{v}' is not a valid date – read as {last:02d}/{mo:02d}/{y}.")
            d = last
        return date(y, mo, d)
    if warnings is not None:
        warnings.append(f"{where}: could not read date '{v}'.")
    return None


def find_sheet(wb, *keywords):
    for ws in wb.worksheets:
        t = ws.title.lower()
        if all(k in t for k in keywords):
            return ws
    return None


def name_from_filename(filename: str) -> str:
    stem = Path(filename).stem
    stem = re.sub(r"\s*-\s*copy.*$", "", stem, flags=re.I)
    stem = re.split(r"[_\s]+summary", stem, flags=re.I)[0]
    stem = stem.replace("_", " ").strip()
    return stem or Path(filename).stem


# --------------------------------------------------------------------------
# Data container
# --------------------------------------------------------------------------
@dataclass
class Project:
    name: str
    code: str
    filename: str
    lead_partner: str = ""
    periods: pd.DataFrame = field(default_factory=pd.DataFrame)   # one row per period
    lines: pd.DataFrame = field(default_factory=pd.DataFrame)     # period x category
    category_budget: pd.DataFrame = field(default_factory=pd.DataFrame)
    staff: pd.DataFrame = field(default_factory=pd.DataFrame)
    spend_check: pd.DataFrame = field(default_factory=pd.DataFrame)
    warnings: list[str] = field(default_factory=list)

    @property
    def label(self) -> str:
        return f"{self.name} ({self.code})" if self.code else self.name


# --------------------------------------------------------------------------
# Main entry point
# --------------------------------------------------------------------------
def parse_project(data: bytes, filename: str) -> Project:
    wb = load_workbook(io.BytesIO(data), data_only=True)
    warnings: list[str] = []

    bud = find_sheet(wb, "budget") or wb.worksheets[0]
    code, lead = "", ""

    # ---- project code & lead partner (search everywhere) -----------------
    for ws in wb.worksheets:
        for row in ws.iter_rows(max_row=min(ws.max_row, 40)):
            for c in row:
                v = c.value
                if isinstance(v, str):
                    if not code and ws is bud and CODE_RE.match(v):
                        code = v.strip().upper()
                    if not lead and "lead partner" in v.lower():
                        after = v.split(":", 1)[1].strip() if ":" in v else ""
                        if not after:
                            right = ws.cell(row=c.row, column=c.column + 1).value
                            below = ws.cell(row=c.row + 1, column=c.column).value
                            after = str(right or below or "").strip()
                        lead = after

    # ---- locate header rows on the budget sheet --------------------------
    rows = [list(r) for r in bud.iter_rows(values_only=True)]
    header_r = next(
        (i for i, r in enumerate(rows[:15]) if any(isinstance(v, str) and "cost category" in v.lower() for v in r)),
        None,
    )
    if header_r is None:
        raise ValueError("Could not find the 'Cost Category' header row on the budget sheet.")
    header = rows[header_r]
    cat_col = next(i for i, v in enumerate(header) if isinstance(v, str) and "cost category" in v.lower())

    period_row_i = next(
        (i for i in range(header_r) if sum(1 for v in rows[i] if isinstance(v, str) and PERIOD_RE.match(v)) >= 2),
        None,
    )
    if period_row_i is None:
        raise ValueError("Could not find the P1, P2, … period labels on the budget sheet.")
    period_cols = [
        (int(PERIOD_RE.match(v).group(1)), i)
        for i, v in enumerate(rows[period_row_i])
        if isinstance(v, str) and PERIOD_RE.match(v)
    ]
    first_pcol = min(c for _, c in period_cols)
    total_budget_col = next(
        (i for i, v in enumerate(header[:first_pcol]) if isinstance(v, str) and v.strip().lower() == "current budget"),
        None,
    )

    # actual vs forecast from the header label of the 2nd column in each pair
    kinds = {}
    for p, c in period_cols:
        lab = str(header[c + 1] if c + 1 < len(header) else "").lower()
        kinds[p] = "Actual" if any(w in lab for w in ACTUAL_WORDS) else "Forecast"

    # ---- cost lines ------------------------------------------------------
    line_recs, cat_budget = [], {}
    for r in rows[header_r + 1:]:
        lab = r[cat_col] if cat_col < len(r) else None
        if isinstance(lab, str) and "partner total" in lab.lower():
            break
        cat = map_category(lab) if lab else None
        if not cat:
            continue
        cb = to_num(r[total_budget_col]) if total_budget_col is not None else 0.0
        per_sum = 0.0
        for p, c in period_cols:
            b = to_num(r[c] if c < len(r) else None)
            v = to_num(r[c + 1] if c + 1 < len(r) else None)
            per_sum += b
            line_recs.append({"period_no": p, "category": cat, "budget": b, "value": v, "kind": kinds[p]})
        if total_budget_col is None:
            cb = per_sum
        elif abs(cb - per_sum) > 1:
            warnings.append(
                f"{cat}: total current budget €{cb:,.2f} ≠ sum of period budgets €{per_sum:,.2f}."
            )
        cat_budget[cat] = cat_budget.get(cat, 0.0) + cb

    lines = pd.DataFrame(line_recs)
    if lines.empty:
        raise ValueError("No cost category rows found on the budget sheet.")
    lines = lines.groupby(["period_no", "category", "kind"], as_index=False)[["budget", "value"]].sum()
    lines["actual"] = lines["value"].where(lines["kind"] == "Actual", 0.0)
    lines["forecast"] = lines["value"].where(lines["kind"] == "Forecast", 0.0)
    lines["period"] = "P" + lines["period_no"].astype(str)

    # ---- reporting periods ----------------------------------------------
    rp = find_sheet(wb, "reporting") or find_sheet(wb, "period")
    prec = []
    if rp is not None:
        for r in rp.iter_rows(min_row=2, values_only=True):
            if not r or not isinstance(r[0], str) or not PERIOD_RE.match(r[0]):
                continue
            p = int(PERIOD_RE.match(r[0]).group(1))
            prec.append({
                "period_no": p,
                "start": to_date(r[1] if len(r) > 1 else None, warnings, f"Reporting periods P{p} start"),
                "finish": to_date(r[2] if len(r) > 2 else None, warnings, f"Reporting periods P{p} finish"),
                "status": str(r[3]).strip() if len(r) > 3 and not is_blank(r[3]) else "",
            })
    periods = pd.DataFrame(prec)
    all_p = sorted(set(kinds) | set(periods["period_no"] if not periods.empty else []))
    periods = (
        pd.DataFrame({"period_no": all_p})
        .merge(periods, on="period_no", how="left") if not periods.empty
        else pd.DataFrame({"period_no": all_p, "start": None, "finish": None, "status": ""})
    )
    periods["status"] = periods["status"].fillna("")
    periods = periods.sort_values("period_no").reset_index(drop=True)
    # fill missing finish from next start
    for i in range(len(periods)):
        if pd.isna(periods.at[i, "finish"]) and i + 1 < len(periods) and not pd.isna(periods.at[i + 1, "start"]):
            periods.at[i, "finish"] = periods.at[i + 1, "start"] - timedelta(days=1)
            warnings.append(f"P{periods.at[i, 'period_no']}: finish date missing – taken as day before next start.")
    periods["period"] = "P" + periods["period_no"].astype(str)
    periods["kind"] = periods["period_no"].map(kinds).fillna("Forecast")

    reported = periods[periods["status"].str.lower().isin(COMPLETED_STATUSES)]
    last_done = int(reported["period_no"].max()) if not reported.empty else 0
    periods["completed"] = periods["period_no"] <= last_done

    def state(row):
        s = row["status"].lower()
        if s == "verified":
            return "Verified"
        if s == "draft":
            return "Draft"
        if row["completed"]:
            return "No claim"
        return "Remaining"

    periods["state"] = periods.apply(state, axis=1)

    # consistency check: actual columns vs reporting status
    last_actual = max((p for p, k in kinds.items() if k == "Actual"), default=0)
    if last_actual != last_done:
        warnings.append(
            f"Budget sheet has actuals up to P{last_actual} but reporting periods show the last "
            f"Verified/Draft period as P{last_done}."
        )
    per_tot = lines.groupby("period")[["budget", "forecast"]].sum()
    for per_lab, r in per_tot[(per_tot["forecast"] > 0) & (per_tot["budget"] == 0)].iterrows():
        warnings.append(f"{per_lab}: forecast €{r['forecast']:,.2f} but no budget phased in this period.")

    # ---- staff -----------------------------------------------------------
    staff_ws = find_sheet(wb, "staff")
    staff = []
    if staff_ws is not None:
        hdr = [str(v or "").strip().lower() for v in next(staff_ws.iter_rows(max_row=1, values_only=True))]
        a_col = next((i for i, h in enumerate(hdr) if "allocation" in h), None)
        for r in staff_ws.iter_rows(min_row=2, values_only=True):
            if not r or is_blank(r[0]):
                continue
            alloc = to_num(r[a_col]) if a_col is not None and a_col < len(r) else None
            if alloc is not None and alloc > 1.0001:   # entered as 25 rather than 0.25
                alloc = alloc / 100
            staff.append({"Staff member": str(r[0]).strip(), "Allocation": alloc})
    staff = pd.DataFrame(staff, columns=["Staff member", "Allocation"])

    # ---- travel & overhead actual spend ---------------------------------
    tv = find_sheet(wb, "travel")
    chk = []
    if tv is not None:
        for r in tv.iter_rows(values_only=True):
            if not r or not isinstance(r[0], str) or not PERIOD_RE.match(r[0]):
                continue
            p = int(PERIOD_RE.match(r[0]).group(1))
            chk.append({
                "period_no": p,
                "travel_spend": None if len(r) < 2 or is_blank(r[1]) else to_num(r[1]),
                "overhead_spend": None if len(r) < 4 or is_blank(r[3]) else to_num(r[3]),
            })
    spend = pd.DataFrame(chk, columns=["period_no", "travel_spend", "overhead_spend"])
    for c in ("travel_spend", "overhead_spend"):
        spend[c] = pd.to_numeric(spend[c], errors="coerce")
    claimed = (
        lines[lines["category"].isin(["Travel & Subsistence", "Overheads"])]
        .pivot_table(index="period_no", columns="category", values="actual", aggfunc="sum")
        .rename(columns={"Travel & Subsistence": "travel_claimed", "Overheads": "overhead_claimed"})
        .reset_index()
    )
    for c in ("travel_claimed", "overhead_claimed"):
        if c not in claimed:
            claimed[c] = 0.0
    act_periods = periods.loc[periods["kind"] == "Actual", ["period_no", "period"]]
    spend_check = act_periods.merge(claimed, on="period_no", how="left").merge(spend, on="period_no", how="left")

    cb = (
        pd.Series(cat_budget, name="current_budget").rename_axis("category").reset_index()
    )

    return Project(
        name=name_from_filename(filename),
        code=code,
        filename=filename,
        lead_partner=lead,
        periods=periods,
        lines=lines,
        category_budget=cb,
        staff=staff,
        spend_check=spend_check,
        warnings=warnings,
    )


# --------------------------------------------------------------------------
# Metrics
# --------------------------------------------------------------------------
def rag(gap_pts: float, amber: float, red: float, flag_under: bool) -> str:
    g = gap_pts if not flag_under else abs(gap_pts)
    if not flag_under and g < 0:
        g = 0
    if g > red:
        return "Red"
    if g > amber:
        return "Amber"
    return "Green"


def project_metrics(p: Project, amber=10.0, red=20.0, flag_under=True) -> dict:
    per = p.periods
    done = per[per["completed"]]
    last = done.iloc[-1] if not done.empty else None
    start = per["start"].dropna().min()
    end = per["finish"].dropna().max()

    if last is not None and start and end and not pd.isna(last["finish"]):
        total_days = (end - start).days + 1
        time_pct = 100 * ((last["finish"] - start).days + 1) / total_days if total_days else 0
    else:
        time_pct = 0.0

    L = p.lines
    budget = float(p.category_budget["current_budget"].sum())
    last_no = int(last["period_no"]) if last is not None else 0
    upto = L[L["period_no"] <= last_no]
    actual_td = float(upto["actual"].sum())
    budget_td = float(upto["budget"].sum())
    actual_all = float(L["actual"].sum())
    forecast_rem = float(L["forecast"].sum())
    fac = actual_all + forecast_rem
    budget_pct = 100 * actual_td / budget if budget else 0.0
    gap = budget_pct - time_pct

    status = f"P{last_no} · {last['status']}" if last is not None else "Not started"
    return {
        "Project": p.name,
        "Code": p.code,
        "Lead partner": p.lead_partner or "—",
        "Status": status,
        "Last status": last["status"] if last is not None else "",
        "Completed periods": int(len(done)),
        "Remaining periods": int(len(per) - len(done)),
        "Total periods": int(len(per)),
        "Start": start,
        "End": end,
        "Last completed": f"P{last_no}" if last_no else "—",
        "% time": time_pct,
        "% budget": budget_pct,
        "Gap (pts)": gap,
        "RAG": rag(gap, amber, red, flag_under),
        "Current budget": budget,
        "Budget to date": budget_td,
        "Actual to date": actual_td,
        "Variance to date": budget_td - actual_td,
        "Forecast remaining": forecast_rem,
        "Forecast at completion": fac,
        "Variance at completion": budget - fac,
    }


def category_table(p: Project) -> pd.DataFrame:
    last_no = int(p.periods.loc[p.periods["completed"], "period_no"].max() or 0) if p.periods["completed"].any() else 0
    L = p.lines
    td = L[L["period_no"] <= last_no].groupby("category")[["budget", "actual"]].sum()
    fc = L.groupby("category")[["actual", "forecast"]].sum()
    t = (
        p.category_budget.set_index("category")
        .join(td.rename(columns={"budget": "Budget to date", "actual": "Actual to date"}))
        .join(fc.rename(columns={"actual": "_act_all", "forecast": "Forecast remaining"}))
        .fillna(0.0)
    )
    t = t.rename(columns={"current_budget": "Current budget"})
    t["Variance to date"] = t["Budget to date"] - t["Actual to date"]
    t["Forecast at completion"] = t["_act_all"] + t["Forecast remaining"]
    t["Variance at completion"] = t["Current budget"] - t["Forecast at completion"]
    t["% spent"] = (100 * t["Actual to date"] / t["Current budget"]).where(t["Current budget"] > 0)
    t = t.drop(columns="_act_all")
    t = t[["Current budget", "Budget to date", "Actual to date", "% spent", "Variance to date",
           "Forecast remaining", "Forecast at completion", "Variance at completion"]]
    order = [c for c in CATEGORY_ORDER if c in t.index]
    t = t.loc[order]
    t = t[(t[["Current budget", "Actual to date", "Forecast at completion"]].abs().sum(axis=1)) > 0]
    return t.reset_index().rename(columns={"category": "Category"})


def periods_ending_in_month(p: Project, today: date | None = None) -> pd.DataFrame:
    """Reporting periods whose finish date falls in the same calendar month as today."""
    today = today or date.today()
    fin = pd.to_datetime(p.periods["finish"], errors="coerce")
    mask = (fin.dt.year == today.year) & (fin.dt.month == today.month)
    return p.periods.loc[mask, ["period", "start", "finish", "status"]].reset_index(drop=True)


def due_text(p: Project, today: date | None = None) -> str:
    """e.g. 'P5 ends 30 Sep' – blank when nothing finishes this month."""
    due = periods_ending_in_month(p, today)
    return ", ".join(f"{r['period']} ends {pd.Timestamp(r['finish']):%d %b}" for _, r in due.iterrows())


def next_period_end(p: Project, today: date | None = None) -> pd.Series | None:
    """First reporting period finishing today or later (None when all periods have finished)."""
    today = today or date.today()
    fin = pd.to_datetime(p.periods["finish"], errors="coerce")
    upcoming = p.periods[fin >= pd.Timestamp(today)].assign(_f=fin).sort_values("_f")
    return None if upcoming.empty else upcoming.iloc[0]


def next_period_text(p: Project, today: date | None = None) -> str:
    """e.g. 'P5 ends 31 Dec 2026'."""
    n = next_period_end(p, today)
    if n is None:
        return "all reporting periods have finished"
    return f"{n['period']} ends {pd.Timestamp(n['finish']):%d %b %Y}"


def next_portfolio_period(projects: list[Project], today: date | None = None) -> str:
    """Earliest upcoming period end across the portfolio, e.g. '31 Oct 2026 – Harbour Link (P3)'."""
    nxt = [(pd.Timestamp(n["finish"]), p.name, n["period"])
           for p in projects if (n := next_period_end(p, today)) is not None]
    if not nxt:
        return ""
    first = min(f for f, _, _ in nxt)
    return f"{first:%d %b %Y} – " + ", ".join(f"{name} ({per})" for f, name, per in sorted(nxt) if f == first)
